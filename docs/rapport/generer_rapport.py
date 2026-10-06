#!/usr/bin/env python3
"""
Génération du rapport de PFE Darimmo (Word)
===========================================

    cd docs/rapport
    ..\\..\\venv\\Scripts\\python.exe generer_rapport.py            # génère le .docx
    ..\\..\\venv\\Scripts\\python.exe generer_rapport.py --word     # puis met à jour les champs avec Word et compte les pages

Fichiers :
    moteur_docx.py      mise en page (consignes FSAC)
    donnees_rapport.py  lecture des chiffres dans tests_ia/, docs/audit_frontend.md, n8n_workflows/
    figures_rapport.py  diagrammes (matplotlib) -> figures/
    contenu_1.py        pages préliminaires, introduction, chapitres 1 à 3
    contenu_2.py        chapitres 4 à 6, conclusion, bibliographie, annexes

Sorties : Rapport_PFE_Darimmo.docx et A_COMPLETER.md (dans ce dossier).
Dépendances : python-docx et matplotlib (dans le venv du projet). L'option --word demande Microsoft Word.
Le script n'écrit rien en dehors de docs/rapport/.
"""

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

import contenu_1
import contenu_2
import donnees_rapport
import figures_rapport
from moteur_docx import Rapport

HERE = Path(__file__).resolve().parent
FSAC = HERE.parent / "fsac"
OUT = HERE / "Rapport_PFE_Darimmo.docx"
TODO = HERE / "A_COMPLETER.md"
ETUDIANT = "Amine Labzioui"

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def _find(folder, keyword):
    for path in sorted(folder.glob("*.docx")) if folder.exists() else []:
        if keyword in path.name.lower():
            return path
    return None


def _set_text(paragraph, text):
    runs = [r for r in paragraph.runs if r.text.strip()] or paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in paragraph.runs:
        if r._r is not runs[0]._r:  # paragraph.runs recrée les objets : on compare l'élément XML
            r.text = ""


def fill_cover(R):
    """Remplit la page de garde FSAC (le modèle sert de base au document)."""
    jury = 0
    for p in R.doc.paragraphs:
        t = p.text.strip()
        if t == "Titre":
            _set_text(p, contenu_1.TITRE)
        elif t == "Nom et Prénom étudiant":
            _set_text(p, ETUDIANT)
        elif t.startswith("Encadrant"):
            _set_text(p, "Encadrant : [À COMPLÉTER]")
        elif t.startswith("Soutenu le"):
            _set_text(p, "Soutenu le : [À COMPLÉTER]")
        elif t == "Nom & Prénom":
            jury += 1
            _set_text(p, "[À COMPLÉTER]")
    R.todo += [
        ("À COMPLÉTER", "Page de garde", "Titre du rapport à valider (proposé à partir du contenu du projet)"),
        ("À COMPLÉTER", "Page de garde", "Encadrant : nom et prénom"),
        ("À COMPLÉTER", "Page de garde", "Date de soutenance (le modèle indiquait « X Juillet 2026 »)"),
        ("À COMPLÉTER", "Page de garde", f"Membres du jury ({jury} lignes)"),
    ]


def plain_cover(R):
    for text, size, bold in (
        ("Université Hassan II de Casablanca — Faculté des Sciences Aïn Chock", 14, True),
        ("Master Ingénierie Informatique et Intelligence Artificielle (IIIA)", 14, True),
        ("Projet de fin d'études", 18, True), (contenu_1.TITRE, 20, True), (ETUDIANT, 16, True),
        ("Encadrant : [À COMPLÉTER]", 14, False), ("Soutenu le : [À COMPLÉTER]", 14, False),
        ("Jury : [À COMPLÉTER]", 14, False), ("Année universitaire : [À COMPLÉTER]", 14, False),
    ):
        R._para(text, align=1, size=size, bold=bold, after=18)


def back_page(R, template, resume):
    """Dernière page FSAC : résumé français et mots-clés dans le modèle, puis ajout à la fin du rapport."""
    doc = Document(str(template))
    done = False
    for table in doc.tables:
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                if "Mettez ici" in p.text:
                    _set_text(p, resume["fr"])
                    done = True
                elif p.text.strip().startswith("Mots-clés") and done:
                    _set_text(p, resume["kw_fr"])
            if done:  # texte blanc sur le fond bleu du modèle, taille réduite pour tenir sur la page
                for p in cell.paragraphs:
                    p.alignment = 3  # justifié
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(10.5)
            if done:
                break
        if done:
            break
    src = doc.sections[0]
    sec = R.new_section(numbered=False, margins=(src.top_margin.cm, src.bottom_margin.cm, src.left_margin.cm, src.right_margin.cm))
    # en-tête et pied de page vides rapprochés du bord : le tableau du modèle occupe presque toute la page
    sec.header_distance = sec.footer_distance = Cm(0.5)
    with tempfile.TemporaryDirectory() as tmp:
        filled = Path(tmp) / "postface.docx"
        doc.save(str(filled))
        R.append_document(filled)
    return done


def write_todo(R, pages=None, words=None):
    words = words or {"fr": "?", "en": "?", "ar": "?"}
    lines = ["# Rapport Darimmo — éléments à compléter ou à relire", "",
             "Fichier généré par `generer_rapport.py` en même temps que `Rapport_PFE_Darimmo.docx`.",
             "Dans le document Word, les marqueurs `[À COMPLÉTER …]` et `[CAPTURE : …]` sont surlignés en jaune.", ""]
    groups = [("À COMPLÉTER", "## Informations à compléter"), ("CAPTURE", "## Captures d'écran à insérer"),
              ("À RELIRE", "## Passages à relire"), ("À VÉRIFIER", "## Points à vérifier")]
    for kind, title in groups:
        items = [t for t in R.todo if t[0] == kind]
        if not items:
            continue
        lines += [title, "", "| Chapitre | Élément |", "|---|---|"]
        seen = set()
        for _, chapter, text in items:
            if (chapter, text) not in seen:
                seen.add((chapter, text))
                lines.append(f"| {chapter} | {text} |")
        lines.append("")
    lines += [
        "## Chiffres sensibles à relire avant l'impression", "",
        "Tous les chiffres des chapitres 5 et des résumés sont lus dans les fichiers au moment de la génération. À contrôler à la lecture :", "",
        "- résumés (français, anglais, arabe) : 77 messages, 27/32 puis 30/32, 59 chiffres de l'Agent Marché, 25 sur 31, 9 cas sur 16 ;",
        "- tableau « Résultats de la classification » : `tests_ia/resultats_classification*/classification_metrics.json` ;",
        "- tableaux des tests de bout en bout : `tests_ia/resultats_e2e_run*/e2e_metrics.json` ;",
        "- tableau et figure de l'Agent Marché : `tests_ia/resultats_e2e_marche_run*/e2e_metrics.json`, `fidelite_metrics.json` (grille `annotation_fidelite_finale.csv`) ;",
        "- tableau de synthèse de l'audit : comptage des lignes des tableaux de `docs/audit_frontend.md` ; « testées dans le navigateur » = lignes dont la colonne « Testé » contient « Étudiant » ou « done » ;",
        "- tableau « Réponses de l'Agent Marché vérifiées à la main » : recopié de `docs/notes_agent_marche_et_session.md` (5 lignes dans ce document).", "",
        "## Écarts avec les consignes à connaître", "",
        f"- Les consignes FSAC demandent un résumé d'environ 300 mots par langue ; les résumés générés en font {words['fr']} (français), {words['en']} (anglais) et {words['ar']} (arabe).",
        "- Les consignes FSAC fixent 100 pages au maximum hors annexes ; le rapport vise 60 pages au maximum, à la demande de l'étudiant.",
        "- Les numéros des titres (1.1, 1.2…) sont écrits dans le texte des titres, pas par une liste numérotée de Word.",
        "- Les sources sont indiquées sous chaque figure et chaque tableau (python-docx ne crée pas de notes de bas de page).",
        "- Après toute modification dans Word : sélectionner tout (Ctrl+A) puis F9 pour mettre à jour la table des matières et les listes.", "",
    ]
    if pages:
        lines += ["## Pagination mesurée par Word", "", f"- Pages au total dans le fichier : {pages['total']}",
                  f"- Pages hors page de garde, dernière page et annexes : {pages['hors_annexes']}",
                  f"- Pages d'annexes : {pages['annexes']}", ""]
    lines += ["## Compteurs", "", f"- Figures : {R.n_fig} (dont {sum(1 for t in R.todo if t[0] == 'CAPTURE')} captures à insérer)",
              f"- Tableaux : {R.n_tab}", ""]
    TODO.write_text("\n".join(lines), encoding="utf-8")


WORD_SCRIPT = r"""
$ErrorActionPreference = 'Stop'
$path = '%s'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
  $doc = $word.Documents.Open($path, $false, $false)
  foreach ($pass in 1..2) {
    $doc.Repaginate()
    foreach ($t in $doc.TablesOfContents) { $t.Update() }
    foreach ($t in $doc.TablesOfFigures) { $t.Update() }
    $null = $doc.Fields.Update()
  }
  $total = $doc.ComputeStatistics(2)
  $annexes = 0; $last = 0
  foreach ($p in $doc.Paragraphs) {
    $txt = $p.Range.Text.Trim()
    if ($txt -eq 'Annexes' -and $p.OutlineLevel -eq 1) { $annexes = $p.Range.Information(3) }
  }
  $last = $doc.Sections.Item($doc.Sections.Count).Range.Information(3)
  $doc.Save()
  "PAGES total=$total annexes=$annexes derniere=$last"
} finally {
  if ($doc) { $doc.Close($false) }
  $word.Quit()
}
"""


def update_with_word():
    """Met à jour les champs (table des matières, listes, numéros) et mesure la pagination avec Word."""
    result = subprocess.run(["powershell", "-NoProfile", "-Command", WORD_SCRIPT % str(OUT)],
                            capture_output=True, text=True, encoding="utf-8", errors="replace")
    line = next((l for l in result.stdout.splitlines() if l.startswith("PAGES")), None)
    if not line:
        print("[info] Word n'a pas pu mettre à jour le document :", (result.stderr or result.stdout).strip()[:300])
        return None
    values = dict(part.split("=") for part in line.split()[1:])
    total, annexes, last = int(values["total"]), int(values["annexes"]), int(values["derniere"])
    return {"total": total, "annexes": (last - annexes) if annexes else 0,
            "hors_annexes": (annexes - 2) if annexes else total - 2}


def main():
    ap = argparse.ArgumentParser(description="Génère le rapport de PFE Darimmo (Word)")
    ap.add_argument("--word", action="store_true", help="met à jour les champs et compte les pages avec Microsoft Word")
    ap.add_argument("--out", default=None, help="nom du fichier Word (utile si le rapport est ouvert dans Word)")
    args = ap.parse_args()
    global OUT
    if args.out:
        OUT = HERE / args.out

    D = donnees_rapport.load()
    F = figures_rapport.tous()
    fid, stab = D["fid"]["chiffres"], D["stab"]
    F["marche"] = figures_rapport.marche_fidelite(
        fid["oui"]["n"], fid["partiel"]["n"], fid["non"]["n"],
        {"Même comportement": stab["comportement"], "Même source citée": stab["source"],
         "Mêmes chiffres": stab["chiffres"], "Réponse identique": stab["identique"],
         "Extraits identiques": stab["extraits_identiques"]})

    cover, back = _find(FSAC, "garde"), _find(FSAC, "post-face")
    R = Rapport(cover)
    if cover:
        fill_cover(R)
    else:
        plain_cover(R)

    R.new_section(page_format="lowerRoman", start=1)
    resume = contenu_1.preliminaires(R, D)

    R.new_section(page_format="decimal", start=1)
    contenu_1.introduction(R, D)
    contenu_1.chapitre_1(R, D)
    contenu_1.chapitre_2(R, D)
    contenu_1.chapitre_3(R, D, F)
    contenu_2.chapitre_4(R, D)
    contenu_2.chapitre_5(R, D, F)
    contenu_2.chapitre_6(R, D)
    contenu_2.conclusion(R, D)
    contenu_2.bibliographie(R)
    contenu_2.annexes(R, D)

    if back:
        if not back_page(R, back, resume):
            R.todo.append(("À COMPLÉTER", "Dernière page", "Résumé et mots-clés non insérés automatiquement dans le modèle FSAC"))
    else:
        R.todo.append(("À COMPLÉTER", "Dernière page", "Modèle de dernière page FSAC absent de docs/fsac"))

    R.save(OUT)
    pages = update_with_word() if args.word else None
    write_todo(R, pages, {k: len(resume[k].split()) for k in ("fr", "en", "ar")})

    words = sum(len(p.text.split()) for p in R.doc.paragraphs)
    print(f"Rapport : {OUT}")
    print(f"Figures : {R.n_fig} | Tableaux : {R.n_tab} | Mots (paragraphes) : {words}")
    if pages:
        print(f"Pages (Word) : {pages['total']} au total | hors page de garde, dernière page et annexes : "
              f"{pages['hors_annexes']} | annexes : {pages['annexes']}")
    print(f"À compléter : {TODO} ({len(R.todo)} éléments)")


if __name__ == "__main__":
    main()
