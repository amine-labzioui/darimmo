"""
Moteur de mise en page du rapport (python-docx), selon les consignes FSAC :
A4, Times New Roman 12 pt, interligne 1,5, texte justifié, marges 2,5 cm (3 cm à gauche),
titres de chapitres 18 pt seuls sur une page, titres de paragraphes 13 pt,
pages préliminaires en chiffres romains, corps en chiffres arabes.

Utilisé par `generer_rapport.py`.
"""

import copy
import io
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


class Rapport:
    def __init__(self, cover_template=None):
        self.doc = Document(str(cover_template)) if cover_template else Document()
        self.has_cover = bool(cover_template)
        self.n_fig = 0
        self.n_tab = 0
        self.todo = []          # (type, chapitre, texte) pour A_COMPLETER.md
        self.chapter = "Pages préliminaires"
        self.figures = []       # (numéro, légende)
        self.tables = []
        self._styles()

    # ------------------------------------------------------------------ styles
    def _font(self, style, size, bold=False, italic=False):
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = BLACK
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), FONT)

    def _ensure_style(self, name, outline_level=None):
        """Crée un style intégré de Word s'il manque dans le modèle (titres, légende)."""
        st = self.doc.styles
        names = [s.name for s in st]
        if name not in names:
            style = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            style.base_style = st["Normal"]
            style.quick_style = True
            if outline_level is not None:
                lvl = OxmlElement("w:outlineLvl")
                lvl.set(qn("w:val"), str(outline_level))
                style.element.get_or_add_pPr().append(lvl)
                nxt = OxmlElement("w:next")
                nxt.set(qn("w:val"), "Normal")
                style.element.insert(2, nxt)
        return st[name]

    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        self._font(normal, 12)
        for level, name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
            self._ensure_style(name, level)
        self._ensure_style("Caption")
        for name, size, before, after in (("Heading 1", 18, 0, 18), ("Heading 2", 13, 14, 8), ("Heading 3", 12, 10, 6)):
            h = st[name]
            self._font(h, size, bold=True)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.15
        st["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._font(st["Caption"], 10, italic=True)

    # ------------------------------------------------------------------ sections
    def new_section(self, page_format=None, start=None, margins=(2.5, 2.5, 3.0, 2.5), numbered=True):
        """Nouvelle section (page suivante) avec les marges FSAC et son propre pied de page."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(margins[0]), Cm(margins[1])
        sec.left_margin, sec.right_margin = Cm(margins[2]), Cm(margins[3])
        sec.footer.is_linked_to_previous = False
        sec.header.is_linked_to_previous = False
        for p in list(sec.footer.paragraphs):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
        pg = sec._sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sec._sectPr.append(pg)
        if page_format:
            pg.set(qn("w:fmt"), page_format)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        if numbered:
            p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._field(p, "PAGE", "1", size=10)
        return sec

    # ------------------------------------------------------------------ champs Word
    def _field(self, paragraph, instruction, cached, size=None, bold=None, italic=None):
        def run_with(child):
            r = paragraph.add_run()
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if italic is not None:
                r.font.italic = italic
            r._element.append(child)
            return r

        for kind in ("begin", None, "separate", "text", "end"):
            if kind is None:
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = f" {instruction} "
                run_with(instr)
            elif kind == "text":
                r = paragraph.add_run(cached)
                if size:
                    r.font.size = Pt(size)
                if bold is not None:
                    r.font.bold = bold
                if italic is not None:
                    r.font.italic = italic
            else:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                if kind == "begin" and instruction.startswith("TOC"):
                    fc.set(qn("w:dirty"), "true")
                run_with(fc)

    # ------------------------------------------------------------------ texte
    def _para(self, text="", style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, bold=False, italic=False,
              spacing=1.5, after=6, before=0, keep=False):
        p = self.doc.add_paragraph(style=style)
        p.alignment = align
        pf = p.paragraph_format
        pf.line_spacing = spacing
        pf.space_after = Pt(after)
        pf.space_before = Pt(before)
        pf.keep_with_next = keep
        if text:
            self._runs(p, text, size, bold, italic)
        return p

    def _runs(self, p, text, size=None, bold=False, italic=False):
        """Texte avec **gras** et `code` en ligne ; les marqueurs [À COMPLÉTER] sont surlignés."""
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[À COMPLÉTER[^\]]*\])", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt((size or 12) - 1.5)
            elif part.startswith("[À COMPLÉTER"):
                r = p.add_run(part)
                r.font.highlight_color = 7  # jaune
                self.todo.append(("À COMPLÉTER", self.chapter, part))
            else:
                r = p.add_run(part)
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
            if bold:
                r.bold = True
            if italic:
                r.italic = True

    def p(self, *texts):
        for t in texts:
            self._para(t)

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.add_run("–  ")
            self._runs(p, it)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def front_title(self, text, toc=True):
        """Titre d'une page préliminaire (non numéroté)."""
        self.chapter = text
        if toc:
            self.doc.add_paragraph(text, style="Heading 1")
        else:
            self._para(text, align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, after=18)

    def chapter_page(self, number, title):
        """Titre de chapitre seul, centré au milieu d'une page (consigne FSAC n° 7)."""
        self.chapter = f"Chapitre {number} — {title}"
        self.page_break()
        for _ in range(11):  # lignes vides : place le titre au milieu de la page
            self._para("", after=0, spacing=1.5)
        h = self.doc.add_paragraph(style="Heading 1")
        h.add_run(f"Chapitre {number}")
        h.add_run().add_break()
        h.add_run(title)
        self.page_break()

    def h2(self, text):
        self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        self.doc.add_paragraph(text, style="Heading 3")

    # ------------------------------------------------------------------ figures et tableaux
    def _caption(self, label, number, text, keep=False):
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = keep
        r = p.add_run(f"{label} ")
        r.bold = True
        self._field(p, f"SEQ {label} \\* ARABIC", str(number), bold=True)
        p.add_run(f" : {text}")
        return p

    def _source(self, text):
        if text:
            self._para(f"Source : {text}", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True, spacing=1.0, after=10)

    def figure(self, path, caption, source=None, width=15.0):
        self.n_fig += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(path), width=Cm(width))
        self._caption("Figure", self.n_fig, caption, keep=bool(source))
        self._source(source or "réalisé par l'auteur")
        self.figures.append((self.n_fig, caption))
        return self.n_fig

    def capture(self, description, caption):
        """Emplacement d'une capture d'écran à insérer par l'étudiant."""
        self.n_fig += 1
        marker = f"[CAPTURE : {description}]"
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Cm(14)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(28)
        cp.paragraph_format.space_after = Pt(28)
        r = cp.add_run(marker)
        r.font.size = Pt(10)
        r.font.highlight_color = 7
        self._para("", after=0, spacing=1.0)
        self._caption("Figure", self.n_fig, caption)
        self._source("capture d'écran de l'application")
        self.todo.append(("CAPTURE", self.chapter, f"Figure {self.n_fig} — {description}"))
        self.figures.append((self.n_fig, caption))
        return self.n_fig

    def table(self, headers, rows, caption, source=None, widths=None, size=10):
        self.n_tab += 1
        self._caption("Tableau", self.n_tab, caption, keep=True)
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            self._cell(t.rows[0].cells[i], h, size, bold=True, fill="E6F4EE")
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                self._cell(cells[i], str(val), size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self._para("", after=0, spacing=1.0, size=4)
        self._source(source)
        self.tables.append((self.n_tab, caption))
        return self.n_tab

    def _cell(self, cell, text, size, bold=False, fill=None):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(1)
        self._runs(p, text, size=size, bold=bold)
        if fill:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            cell._element.get_or_add_tcPr().append(shd)

    @staticmethod
    def _repeat_header(row):
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(el)

    def code(self, text, size=8.5):
        for line in text.split("\n"):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(size)
        self._para("", after=4, spacing=1.0)

    def arabic(self, text, size=13):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        ppr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
        r = p.add_run(text)
        rpr = r._element.get_or_add_rPr()
        rpr.append(OxmlElement("w:rtl"))
        sz = OxmlElement("w:szCs")
        sz.set(qn("w:val"), str(size * 2))
        rpr.append(sz)
        return p

    # ------------------------------------------------------------------ tables des matières
    def toc(self, instruction, placeholder):
        p = self.doc.add_paragraph()
        self._field(p, instruction, placeholder)

    # ------------------------------------------------------------------ fusion d'un autre document
    def append_document(self, path):
        """Copie le corps d'un document Word (dernière page FSAC) à la fin, images comprises."""
        other = Document(str(path))
        body = self.doc.element.body
        main_sect = body.find(qn("w:sectPr"))
        for el in other.element.body:
            if el.tag == qn("w:sectPr"):
                continue
            new = copy.deepcopy(el)
            if new.tag == qn("w:p") and not "".join(t.text or "" for t in new.iter(qn("w:t"))).strip()                     and new.find(".//" + qn("w:drawing")) is None:
                # paragraphe vide du modèle : réduit à 1 pt pour que la page garde sa hauteur d'origine
                ppr = new.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr")
                    new.insert(0, ppr)
                for old in ppr.findall(qn("w:spacing")):
                    ppr.remove(old)
                spacing = OxmlElement("w:spacing")
                for key, val in (("w:before", "0"), ("w:after", "0"), ("w:line", "20"), ("w:lineRule", "exact")):
                    spacing.set(qn(key), val)
                ppr.append(spacing)
            for blip in new.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                if rid:
                    # l'image est recréée dans ce document (un simple lien garderait le même nom de fichier interne)
                    blob = other.part.related_parts[rid].blob
                    new_rid, _ = self.doc.part.get_or_add_image(io.BytesIO(blob))
                    blip.set(qn("r:embed"), new_rid)
            if main_sect is not None:
                main_sect.addprevious(new)
            else:
                body.append(new)
        return other

    def save(self, path):
        self.doc.save(str(path))
